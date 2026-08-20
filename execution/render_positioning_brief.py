#!/usr/bin/env python3
"""Render a positioning brief to DOCX.

⛔ APPROVAL-GATED (BR-5, BR-19). This tool writes to the working tree only. It never
commits and never pushes. The sequence is: present the analysis → the user reviews and
challenges it → the user approves → only then is saving discussed.

⛔ BR-3 / P3 / F6: a novelty claim is never derived from absence of evidence. The brief
carries a mandatory coverage statement naming what was searched and what was NOT, and
refuses to emit a gap claim that the coverage cannot support.

Verify: python3 execution/tests/test_render_positioning_brief.py
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "execution"))
from validate_registry import validate_record  # noqa: E402

BRIEF_DIR = REPO_ROOT / "positioning-briefs"

# Categories that speak directly to positioning, in the order a brief should present
# them: what the work stands on, what it stands next to, what could knock it over.
SECTIONS = [
    ("foundational", "Foundational work this builds on"),
    ("methodological_precedent", "Methodological precedents"),
    ("closest_to_work", "Closest existing work"),
    ("novelty_threat", "Threats to the novelty claim"),
    ("challenges_assumption", "Work that challenges an assumption made here"),
    ("needs_citation_support", "Claims still needing literature support"),
]


def by_category(records: list[dict], category: str) -> list[dict]:
    return [r for r in records
            if category in ((r.get("screening") or {}).get("categories") or [])]


def citation(record: dict) -> tuple[str, bool]:
    """Return (display string, citable). Non-citable records are labelled, never dropped."""
    bib = record.get("bibliographic") or {}
    ids = record.get("identifiers") or {}
    authors = bib.get("authors") or []
    who = (authors[0].split(",")[0] + (" et al." if len(authors) > 1 else "")) if authors else "—"
    doi = ids.get("doi")
    text = f"{who} ({bib.get('year') or 'n.d.'}). {bib.get('title') or '—'}."
    if bib.get("venue"):
        text += f" {bib['venue']}."
    if doi:
        text += f" doi:{doi}"
    return text, not validate_record(record)[1]


def write_docx(idea: str, records: list[dict], path: Path, anchor: str | None,
               sources_searched: list[str], not_searched: list[str]) -> dict:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(f"Positioning brief — {idea}", level=0)

    head = doc.add_paragraph()
    head.add_run("Status: DRAFT — not approved, not committed.").bold = True
    head.add_run(" This brief is intellectual output and is approval-gated (BR-5). "
                 "Review and challenge it before any part of it is treated as settled.")

    if anchor:
        p = doc.add_paragraph()
        p.add_run("Positioned against: ").bold = True
        p.add_run(anchor)

    # --- The coverage statement comes FIRST, deliberately -----------------------
    # A reader must know the limits of the search before reading any conclusion drawn
    # from it. Putting this at the end would let the conclusions land unqualified.
    doc.add_heading("Coverage of this search", level=1)
    cov = doc.add_paragraph()
    cov.add_run("Searched: ").bold = True
    cov.add_run(", ".join(sources_searched) if sources_searched else "not recorded")
    if not_searched:
        c2 = doc.add_paragraph()
        c2.add_run("NOT searched: ").bold = True
        c2.add_run(", ".join(not_searched))
    warn = doc.add_paragraph()
    wr = warn.add_run(
        "Absence of an identical paper in the sources above is NOT evidence of a gap "
        "(BR-3, P3). Any novelty claim below is bounded by this coverage, and adjacent "
        "literatures using different terminology may not be represented.")
    wr.italic = True
    wr.font.size = Pt(9)

    counts = {}
    noncitable_total = 0
    for key, heading in SECTIONS:
        group = by_category(records, key)
        counts[key] = len(group)
        if not group:
            continue
        doc.add_heading(heading, level=1)
        for r in group:
            text, citable = citation(r)
            reason = ((r.get("screening") or {}).get("reason") or "").strip()
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(text)
            if reason:
                p.add_run(f" — {reason}")
            if not citable:
                noncitable_total += 1
                flag = doc.add_paragraph()
                fr = flag.add_run(
                    "    ⚠ Candidate only — surfaced for manual checking, not cited as "
                    "support (§3.0).")
                fr.font.size = Pt(9)
                fr.bold = True

    doc.add_heading("What this brief does not establish", level=1)
    limits = doc.add_paragraph(style="List Bullet")
    limits.add_run(
        "Whether the contribution is strong enough. That is a judgement for the author, "
        "informed by the threats listed above — this document assembles evidence, it "
        "does not rule.")
    doc.add_paragraph(style="List Bullet").add_run(
        f"Anything about the {len(not_searched)} unsearched source(s) named above.")
    if noncitable_total:
        doc.add_paragraph(style="List Bullet").add_run(
            f"Anything resting on the {noncitable_total} candidate-only record(s), whose "
            f"identifiers were not verified or whose metadata is disputed.")

    doc.add_paragraph()
    foot = doc.add_paragraph()
    fr2 = foot.add_run(
        "Generated from state/paper-registry.json. Every citation above is a registry "
        "record; this document adds no bibliographic detail of its own. DRAFT — "
        "approval-gated, never auto-committed (BR-5, BR-19).")
    fr2.font.size = Pt(8)
    fr2.italic = True

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return {"by_category": counts, "noncitable": noncitable_total}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--idea", required=True, help="short name for the idea being positioned")
    ap.add_argument("--registry", default=str(REPO_ROOT / "state" / "paper-registry.json"))
    ap.add_argument("--out", default=None)
    ap.add_argument("--anchor", default=None)
    ap.add_argument("--searched", default="", help="comma-separated sources actually searched")
    ap.add_argument("--not-searched", default="", help="comma-separated sources NOT searched")
    ap.add_argument("--demo", action="store_true")
    args = ap.parse_args()

    if args.demo:
        from render_screening_report import demo_records
        records = demo_records()
    else:
        p = Path(args.registry)
        if not p.exists():
            print(json.dumps({"tool": "render_positioning_brief", "status": "invalid",
                              "detail": f"{p} not found"}))
            return 4
        records = json.loads(p.read_text()).get("records", [])

    searched = [s.strip() for s in args.searched.split(",") if s.strip()]
    not_searched = [s.strip() for s in args.not_searched.split(",") if s.strip()]

    # F6 guard: a positioning brief with no recorded coverage cannot bound any claim
    # it makes, so it is refused rather than emitted with an empty caveat.
    if not searched:
        print(json.dumps({
            "tool": "render_positioning_brief", "status": "missing",
            "detail": "No sources recorded as searched. A positioning brief states what "
                      "was and was not covered (BR-3, P3, F6); without that its "
                      "conclusions cannot be bounded. Pass --searched."}))
        return 3

    slug = "".join(c if c.isalnum() or c == "-" else "-" for c in args.idea.lower())[:50].strip("-")
    out = Path(args.out) if args.out else BRIEF_DIR / f"{slug}-{date.today().isoformat()}.docx"
    stats = write_docx(args.idea, records, out, args.anchor, searched, not_searched)
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")

    print(json.dumps({
        "tool": "render_positioning_brief", "status": "green",
        "detail": f"DRAFT written to the working tree. Not committed — approval-gated "
                  f"(BR-5, BR-19).",
        "evidence": {"path": str(out), "records": len(records), **stats,
                     "sources_searched": searched, "sources_not_searched": not_searched,
                     "committed": False},
        "checked_at": now}))
    print(f"[OK] DRAFT {out} — review before this is saved anywhere (BR-5)", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
