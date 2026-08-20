#!/usr/bin/env python3
"""Extract an anchor document to plain text. See SOP-003 section 4.2.

The anchor is what relevance is judged AGAINST (BR-10). Ranking papers by broad terms
like "CAR-T" or "supply chain" instead of against this text is failure mode F5.

Verify: python3 execution/tests/test_screening_pipeline.py
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "execution" / "probes"))
from _common import GREEN, INVALID, MISSING, emit  # noqa: E402

TOOL = "extract_anchor"


def extract(path: Path) -> tuple[str, dict]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        tables = 0
        for t in doc.tables:
            tables += 1
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(x for x in parts if x.strip())
        return text, {"format": "docx", "paragraphs": len(doc.paragraphs), "tables": tables}
    if suffix in (".md", ".txt", ".tex"):
        text = path.read_text(encoding="utf-8", errors="replace")
        return text, {"format": suffix.lstrip(".")}
    raise ValueError(f"Unsupported anchor format {suffix!r}. Use .docx, .md, .txt or .tex.")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--path", required=True, help="the anchor document the user named (BR-9)")
    ap.add_argument("--out", default=str(REPO_ROOT / ".tmp" / "anchor.txt"))
    args = ap.parse_args()

    src = Path(args.path)
    if not src.exists():
        return emit(TOOL, MISSING, f"Anchor document not found: {src}. Screening cannot "
                                   f"proceed without it — this is a schema error, not a "
                                   f"degraded run (BR-10).", {"path": str(src)})
    try:
        text, meta = extract(src)
    except Exception as exc:  # noqa: BLE001
        return emit(TOOL, INVALID, str(exc), {"path": str(src)})

    if len(text.strip()) < 200:
        return emit(TOOL, INVALID,
                    f"Anchor extracted to only {len(text.strip())} characters. Too thin to "
                    f"judge relevance against — check the file is the right one.",
                    {"path": str(src), "chars": len(text.strip())})

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")

    return emit(TOOL, GREEN, f"Anchor extracted: {len(text)} characters.",
                {"source": str(src), "path": str(out), "chars": len(text), **meta,
                 "extracted_at": datetime.now(timezone.utc).isoformat(timespec="seconds")})


if __name__ == "__main__":
    sys.exit(main())
