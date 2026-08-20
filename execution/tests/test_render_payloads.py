#!/usr/bin/env python3
"""Tests for the BUILD note, digest and positioning brief renderers (SOP-004 §7).

Each suite targets the rule that payload exists to honour: BR-1/P4 for the note,
BR-18 for the digest, BR-3/P3/F6 and BR-5 for the brief.

Verify: python3 execution/tests/test_render_payloads.py
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
EXEC = REPO_ROOT / "execution"
sys.path.insert(0, str(EXEC))

from render_build_note import build_note, note_path, parse_note, slugify  # noqa: E402
from render_digest import qualifying  # noqa: E402
from render_positioning_brief import by_category, citation  # noqa: E402
from render_screening_report import demo_records  # noqa: E402

FAILURES: list[str] = []
TMP = Path(tempfile.mkdtemp())


def check(name: str, condition: bool, detail: str = "") -> None:
    print(f"  {'✅' if condition else '❌'} {name}{'' if condition else '  ' + detail}")
    if not condition:
        FAILURES.append(name)


def run(script: str, *args) -> tuple[int, dict]:
    proc = subprocess.run([sys.executable, str(EXEC / script), *args],
                          capture_output=True, text=True)
    try:
        payload = json.loads((proc.stdout or "").strip().splitlines()[-1])
    except Exception:
        payload = {}
    return proc.returncode, payload


# ───────────────────────── BUILD notes ─────────────────────────
print("\n  BUILD note — BR-1 / P4: the reader does the work\n")

note = build_note("A paper", {}, "2026-08-19T00:00:00+00:00")
check("all five BUILD stages appear", all(f"## {k}:" in note for k in "BUILD"))
check("stage headings match the skill's own wording",
      "## B: Reading purposes" in note and "## D: What I learned" in note)
check("incomplete stages are explicit placeholders, never filled in",
      note.count("_(not yet completed)_") == 5)
check("the L stage carries its verdict hint", "Aligned / Diverges / Unresolved" in note)
check("the file states it contains no summary", "no summary of the paper" in note)

filled = build_note("A paper", {"B": "my purposes", "I": "my words"},
                    "2026-08-19T00:00:00+00:00")
check("supplied stages render the reader's own text",
      "my purposes" in filled and "my words" in filled)
check("unsupplied stages stay placeholders", filled.count("_(not yet completed)_") == 3)

title, stages = parse_note(filled)
check("an existing note round-trips its title", title == "A paper", title)
check("an existing note round-trips completed stages only",
      set(stages) == {"B", "I"}, str(set(stages)))

p = TMP / "note.md"
rc, _ = run("render_build_note.py", "--title", "Round trip", "--out", str(p))
rc, _ = run("render_build_note.py", "--title", "Round trip", "--out", str(p),
            "--stage", "B", "--content", "first stage")
rc, out = run("render_build_note.py", "--title", "Round trip", "--out", str(p),
              "--stage", "U", "--content", "second stage")
body = p.read_text()
check("incremental writes preserve earlier stages",
      "first stage" in body and "second stage" in body)
check("progress is reported honestly",
      out["evidence"]["stages_complete"] == ["B", "U"], str(out.get("evidence")))

rc, out = run("render_build_note.py", "--title", "X", "--out", str(TMP / "x.md"), "--stage", "I")
check("a stage with no content is refused, never invented",
      rc == 3 and "never writes stage content" in out.get("detail", ""), str(rc))

check("slug is filesystem-safe", slugify("Optimisation: CAR-T & scheduling!") ==
      "optimisation-car-t-scheduling")
check("note path follows YYYY-MM-DD-short-title.md",
      note_path("A Title").name.endswith("-a-title.md"))

# ───────────────────────── Digest ─────────────────────────
print("\n  Digest — BR-18: an empty digest is a success\n")

recs = demo_records()
q = qualifying(recs)
check("only 'relevant' verdicts qualify",
      all(r["screening"]["verdict"] == "relevant" for r in q), str(len(q)))
check("'uncertain' does not qualify for a digest",
      not any(r["screening"]["verdict"] == "uncertain" for r in q))
check("a relevant record with no reason does not qualify",
      qualifying([{"screening": {"verdict": "relevant", "reason": "  "}}]) == [])

empty_reg = TMP / "empty.json"
empty_reg.write_text(json.dumps({"schema_version": "1.0.0", "records": []}))
target = TMP / "must-not-exist.docx"
rc, out = run("render_digest.py", "--registry", str(empty_reg), "--out", str(target))
check("zero qualifying papers exits 0 — it is a success, not a failure", rc == 0, str(rc))
check("zero qualifying papers writes NO file", not target.exists())
check("the outcome is recorded as no_qualifying_papers",
      out["evidence"]["outcome"] == "no_qualifying_papers")
check("artifact is explicitly null", out["evidence"]["artifact"] is None)
check("the message says an empty digest is correct", "correct outcome" in out["detail"])

full = TMP / "digest.docx"
rc, out = run("render_digest.py", "--demo", "--out", str(full), "--period", "2026-08-24")
check("a digest with qualifying papers is written", rc == 0 and full.exists())
check("candidates_considered and papers_included are both reported",
      out["evidence"]["candidates_considered"] == 6
      and out["evidence"]["papers_included"] == len(q), str(out.get("evidence")))

from docx import Document  # noqa: E402
text = "\n".join(p.text for p in Document(str(full)).paragraphs)
check("each entry states why it may matter", "Why it may matter:" in text)
check("the digest disclaims being a summary", "not a literature summary" in text)
check("evidence inspected is carried into the digest", "Evidence inspected:" in text)
check("non-citable records are flagged in the digest", "Candidate only" in text)

# ───────────────────────── Positioning brief ─────────────────────────
print("\n  Positioning brief — BR-3 / P3 / F6 and BR-5\n")

rc, out = run("render_positioning_brief.py", "--idea", "x", "--demo",
              "--out", str(TMP / "b.docx"))
check("a brief with no recorded coverage is refused", rc == 3, str(rc))
check("the refusal names the rule", "BR-3, P3, F6" in out.get("detail", ""))
check("no file is written when refused", not (TMP / "b.docx").exists())

brief = TMP / "brief.docx"
rc, out = run("render_positioning_brief.py", "--idea", "survival-aware scheduling",
              "--demo", "--out", str(brief),
              "--searched", "Zotero, OpenAlex", "--not-searched", "Scopus, Web of Science")
check("a brief with coverage is written", rc == 0 and brief.exists())
check("the result records that it was NOT committed",
      out["evidence"]["committed"] is False)
check("the message says approval-gated", "approval-gated" in out["detail"])

btext = "\n".join(p.text for p in Document(str(brief)).paragraphs)
check("the brief is marked DRAFT at the top", "Status: DRAFT" in btext)
check("coverage appears before any positioning section",
      btext.index("Coverage of this search") < btext.index("Closest existing work"))
check("unsearched sources are named", "Scopus" in btext and "Web of Science" in btext)
check("absence-is-not-a-gap is stated explicitly",
      "NOT evidence of a gap" in btext)
check("adjacent-terminology limitation is stated",
      "different terminology" in btext)
check("threats to novelty get their own section", "Threats to the novelty claim" in btext)
check("the brief declines to rule on contribution strength",
      "does not rule" in btext)
check("non-citable records are surfaced but not cited as support",
      "Candidate only" in btext and "not cited as support" in btext)

cited, citable = citation(demo_records()[0])
check("a citation is built only from registry fields", "doi:10.1000/demo.full" in cited)
check("a verified record is citable", citable)
_, uncitable = citation([r for r in demo_records() if not r["identifier_verified_against_source"]][0])
check("an unverified record is not citable", not uncitable)
check("category grouping selects on membership",
      len(by_category(demo_records(), "closest_to_work")) == 2)

print()
if FAILURES:
    print(f"  ❌ {len(FAILURES)} failed: {FAILURES}\n")
    sys.exit(1)
print("  ✅ All payload tests passed.\n")
