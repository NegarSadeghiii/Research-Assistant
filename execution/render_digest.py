#!/usr/bin/env python3
"""Render a monitoring digest to DOCX — or, correctly, to nothing at all.

⛔ BR-18 is the load-bearing rule here: zero qualifying papers produces NO FILE and a
short status message. An empty digest is a correct and desirable outcome, not a gap to
fill. A system obliged to produce N papers per run will find N papers whether or not
they merit attention — that is F2 arriving on a schedule.

⛔ Not a literature summary. Each entry answers "why might this matter to the current
research", which is a justification, not a précis.

Verify: python3 execution/tests/test_render_digest.py
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "execution"))
from validate_registry import validate_record  # noqa: E402
from render_screening_report import evidence_class  # noqa: E402

DIGEST_DIR = REPO_ROOT / "literature-digests"


def qualifying(records: list[dict]) -> list[dict]:
    """A record qualifies only on a `relevant` verdict with a stated reason.

    Deliberately strict. 'uncertain' does not qualify: a digest exists to say
    "this is worth your attention", and uncertainty cannot support that claim.
    Uncertain records stay in the screening report, where they are visible.
    """
    out = []
    for r in records:
        s = r.get("screening") or {}
        if s.get("verdict") == "relevant" and (s.get("reason") or "").strip():
            out.append(r)
    return out


def write_docx(records: list[dict], path: Path, period: str, anchor: str | None) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(f"Literature digest — {period}", level=0)

    intro = doc.add_paragraph()
    intro.add_run(f"{len(records)} paper(s) met the relevance threshold").bold = True
    if anchor:
        intro.add_run(f", screened against {anchor}")
    intro.add_run(".")

    for r in records:
        bib = r.get("bibliographic") or {}
        s = r.get("screening") or {}
        ids = r.get("identifiers") or {}
        _, blocks = validate_record(r)
        _, ev_label = evidence_class(r)

        doc.add_heading(bib.get("title") or "—", level=2)

        meta = doc.add_paragraph()
        authors = bib.get("authors") or []
        who = authors[0] + (f" et al." if len(authors) > 1 else "") if authors else "—"
        run = meta.add_run(f"{who} · {bib.get('year') or '—'} · {bib.get('venue') or '—'}")
        run.italic = True
        run.font.size = Pt(9)

        why = doc.add_paragraph()
        why.add_run("Why it may matter: ").bold = True
        why.add_run(s.get("reason") or "")

        cats = s.get("categories") or []
        if cats:
            p = doc.add_paragraph()
            r_ = p.add_run("Categories: " + ", ".join(c.replace("_", " ") for c in cats))
            r_.font.size = Pt(9)

        prov = doc.add_paragraph()
        pr = prov.add_run(f"Evidence inspected: {ev_label}. "
                          f"Identifier: {ids.get('doi') or 'none recorded'}.")
        pr.font.size = Pt(9)
        pr.italic = True

        if blocks:
            warn = doc.add_paragraph()
            wr = warn.add_run("⚠ Candidate only — not citable as evidence: "
                              + "; ".join(blocks))
            wr.bold = True
            wr.font.size = Pt(9)

    doc.add_paragraph()
    foot = doc.add_paragraph()
    fr = foot.add_run(
        "Generated from state/paper-registry.json. Each entry states why the paper may "
        "matter to current research; this document is not a literature summary. Papers "
        "judged uncertain or irrelevant are not omitted from the record — they remain "
        "in the screening report for this period.")
    fr.font.size = Pt(8)
    fr.italic = True

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--registry", default=str(REPO_ROOT / "state" / "paper-registry.json"))
    ap.add_argument("--out", default=None)
    ap.add_argument("--period", default=date.today().isoformat())
    ap.add_argument("--anchor", default=None)
    ap.add_argument("--demo", action="store_true")
    args = ap.parse_args()

    if args.demo:
        from render_screening_report import demo_records
        records = demo_records()
    else:
        p = Path(args.registry)
        if not p.exists():
            print(json.dumps({"tool": "render_digest", "status": "invalid",
                              "detail": f"{p} not found"}))
            return 4
        records = json.loads(p.read_text()).get("records", [])

    included = qualifying(records)
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")

    # BR-18. This branch is the whole point of the tool.
    if not included:
        print(json.dumps({
            "tool": "render_digest", "status": "green",
            "detail": "No qualifying papers for this period. No file written — an empty "
                      "digest is a correct outcome, not a failure (BR-18).",
            "evidence": {"candidates_considered": len(records), "papers_included": 0,
                         "artifact": None, "outcome": "no_qualifying_papers"},
            "checked_at": now}))
        print(f"[OK] no qualifying papers ({len(records)} considered) — no file written",
              file=sys.stderr)
        return 0

    out = Path(args.out) if args.out else DIGEST_DIR / f"{args.period}.docx"
    write_docx(included, out, args.period, args.anchor)

    print(json.dumps({
        "tool": "render_digest", "status": "green",
        "detail": f"Digest written with {len(included)} paper(s).",
        "evidence": {"candidates_considered": len(records),
                     "papers_included": len(included), "artifact": str(out),
                     "outcome": "digest_written", "bytes": out.stat().st_size},
        "checked_at": now}))
    print(f"[OK] wrote {out} ({len(included)} papers)", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
